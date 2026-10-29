"""生成 docx 验收样本（sample_reports/星澜科技2025年度简报.docx）

虚构公司 + 写死的数字：既当 docx 解析管线的 smoke 样本（段落/表格/图片三类
chunk 齐全），又当评测锚定语料（evalset M05 的标准出处）。数字与真语料零重合，
不会被年报题误命中。重新运行 = 覆盖重建（内容确定，幂等）。

用法：python -X utf8 tools/make_docx_sample.py
"""

import io
import struct
import zlib
from pathlib import Path

from docx import Document

OUT = Path(__file__).parent.parent / "sample_reports" / "星澜科技2025年度简报.docx"


def _png(width=240, height=120, rgb=(64, 116, 210)) -> bytes:
    """stdlib 手搓纯色 PNG——样本图只为走通「docx 图片 → VLM 描述 → image chunk」
    管线，不值得为它引 Pillow。"""
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))

    def chunk(tag: bytes, data: bytes) -> bytes:
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw))
            + chunk(b"IEND", b""))


def main():
    doc = Document()
    doc.add_heading("星澜科技2025年度经营简报", level=1)
    doc.add_paragraph(
        "星澜科技（虚构演示公司）主营工业智能传感器与边缘计算设备的研发、制造与销售，"
        "产品应用于离散制造产线的状态监测与能耗管理。本简报为 ragbot docx 导入管线的"
        "验收样本，其中数字全部虚构，与任何真实公司无关。")
    doc.add_paragraph(
        "2025年，制造业智能化改造需求持续释放，公司在汽车零部件与3C电子两个行业线"
        "均实现了客户数量与单客户采购额的双增长；海外市场方面，东南亚区域完成本地化"
        "服务网点布局，出口收入占比提升至18.4%。")

    doc.add_heading("主要财务指标", level=2)
    doc.add_paragraph("下表列示公司近三年主要财务数据（单位：元，除标注外，虚构）。")
    # 12 行指标——刻意超过 store.SMALL_TABLE_ROWS(10)，让大表格走「独立 table chunk」
    # 路径（内联小表会跟正文混切成 text chunk，evalset 的表格锚定就失配了）
    rows = [
        ("指标", "2023年度", "2024年度", "2025年度"),
        ("营业收入（元）", "612,345,678.90", "745,678,901.23", "876,543,200.55"),
        ("归属于母公司股东的净利润（元）", "58,901,234.56", "81,234,567.89", "102,345,678.90"),
        ("扣除非经常性损益的净利润（元）", "55,123,456.78", "78,456,789.01", "98,765,432.10"),
        ("经营活动产生的现金流量净额（元）", "71,234,567.89", "120,345,678.90", "158,900,000.00"),
        ("基本每股收益（元/股）", "0.61", "0.84", "1.06"),
        ("加权平均净资产收益率", "9.8%", "12.4%", "14.1%"),
        ("研发投入（元）", "61,234,567.89", "82,345,678.90", "103,432,098.76"),
        ("研发投入占营业收入比例", "10.0%", "11.0%", "11.8%"),
        ("总资产（元）", "890,123,456.78", "1,012,345,678.90", "1,145,678,901.23"),
        ("归属于上市公司股东的净资产（元）", "654,321,098.76", "723,456,789.01", "801,234,567.89"),
        ("资产负债率", "26.5%", "28.5%", "30.0%"),
        ("存货周转率（次）", "2.8", "3.1", "3.3"),
    ]
    table = doc.add_table(rows=len(rows), cols=4)
    table.style = "Table Grid"
    for r, cells in enumerate(rows):
        for c, val in enumerate(cells):
            table.cell(r, c).text = val

    doc.add_paragraph("下图展示了公司近三年营业收入的增长趋势（示意图，虚构数据）。")
    doc.add_picture(io.BytesIO(_png()))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"[OK] {OUT}")


if __name__ == "__main__":
    main()
